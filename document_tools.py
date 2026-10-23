"""Offline document extraction and conservative, copy-only edits."""
from __future__ import annotations

import re
import csv
import html
import io
import zipfile
import shutil
import time
import uuid
import copy
from functools import lru_cache
from enum import Enum
from pathlib import Path


SUPPORTED = {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".htm", ".txt", ".md", ".rtf"}


def is_document(path: str | Path) -> bool:
    return Path(path).suffix.lower() in SUPPORTED


def _docx_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _html_text(source: str) -> str:
    from html.parser import HTMLParser

    class TextReader(HTMLParser):
        def __init__(self):
            super().__init__(convert_charrefs=True)
            self.parts = []
            self.hidden = 0

        def handle_starttag(self, tag, attrs):
            if tag.casefold() in {"script", "style", "noscript"}:
                self.hidden += 1
            elif tag.casefold() in {"br", "p", "div", "li", "tr", "h1", "h2", "h3", "section"}:
                self.parts.append("\n")

        def handle_endtag(self, tag):
            if tag.casefold() in {"script", "style", "noscript"} and self.hidden:
                self.hidden -= 1
            elif tag.casefold() in {"p", "div", "li", "tr", "h1", "h2", "h3", "section"}:
                self.parts.append("\n")

        def handle_data(self, data):
            if not self.hidden:
                self.parts.append(data)

    parser = TextReader()
    parser.feed(source)
    return re.sub(r"[ \t]+\n", "\n", "".join(parser.parts)).strip()


def _ooxml_text(path: Path) -> tuple[str, dict]:
    """Read visible text from PPTX/XLSX using the OOXML package format (stdlib only)."""
    from xml.etree import ElementTree as ET

    rows = []
    with zipfile.ZipFile(path, "r") as package:
        names = set(package.namelist())
        if path.suffix.lower() == ".pptx":
            slide_names = sorted(
                (name for name in names if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)),
                key=lambda value: int(re.search(r"slide(\d+)", value).group(1)),
            )
            for number, name in enumerate(slide_names, start=1):
                root = ET.fromstring(package.read(name))
                paragraphs = []
                for paragraph in root.iter():
                    if paragraph.tag.rsplit("}", 1)[-1] != "p":
                        continue
                    value = "".join(node.text or "" for node in paragraph.iter() if node.tag.rsplit("}", 1)[-1] == "t")
                    if value.strip():
                        paragraphs.append(value)
                if paragraphs:
                    rows.append(f"[Slide {number}]\n" + "\n".join(paragraphs))
            return "\n\n".join(rows), {"slides": len(slide_names), "images": 0}

        workbook_name = "xl/workbook.xml"
        rels_name = "xl/_rels/workbook.xml.rels"
        if workbook_name not in names or rels_name not in names:
            raise ValueError("This XLSX file is missing its workbook index and cannot be read safely.")
        workbook = ET.fromstring(package.read(workbook_name))
        rels = ET.fromstring(package.read(rels_name))
        rel_targets = {
            rel.attrib.get("Id", ""): rel.attrib.get("Target", "")
            for rel in rels.iter()
            if rel.tag.rsplit("}", 1)[-1] == "Relationship"
        }
        shared = []
        if "xl/sharedStrings.xml" in names:
            strings_root = ET.fromstring(package.read("xl/sharedStrings.xml"))
            for item in strings_root.iter():
                if item.tag.rsplit("}", 1)[-1] == "si":
                    shared.append("".join(node.text or "" for node in item.iter() if node.tag.rsplit("}", 1)[-1] == "t"))
        sheets = []
        for sheet in workbook.iter():
            if sheet.tag.rsplit("}", 1)[-1] != "sheet":
                continue
            title = sheet.attrib.get("name", "Sheet")
            rel_id = next((value for key, value in sheet.attrib.items() if key.endswith("}id")), "")
            target = rel_targets.get(rel_id, "")
            if not target:
                continue
            sheet_path = target.lstrip("/") if target.startswith("/") else "xl/" + target
            # Relationship targets may contain ./ or ../ path segments.
            import posixpath
            sheet_path = posixpath.normpath(sheet_path)
            if sheet_path not in names:
                continue
            root = ET.fromstring(package.read(sheet_path))
            sheet_rows = []
            for row in root.iter():
                if row.tag.rsplit("}", 1)[-1] != "row":
                    continue
                values = []
                for cell in row:
                    if cell.tag.rsplit("}", 1)[-1] != "c":
                        continue
                    cell_type = cell.attrib.get("t", "")
                    value = ""
                    if cell_type == "inlineStr":
                        value = "".join(node.text or "" for node in cell.iter() if node.tag.rsplit("}", 1)[-1] == "t")
                    else:
                        raw = next((node.text or "" for node in cell if node.tag.rsplit("}", 1)[-1] == "v"), "")
                        if cell_type == "s" and raw.isdigit() and int(raw) < len(shared):
                            value = shared[int(raw)]
                        else:
                            value = raw
                    values.append(value)
                if any(values):
                    sheet_rows.append("\t".join(values))
            if sheet_rows:
                sheets.append(f"[Sheet: {title}]\n" + "\n".join(sheet_rows))
        return "\n\n".join(sheets), {"sheets": len(sheets), "images": 0}


def _replace_ooxml_text(path: Path, find: str, replace: str) -> int:
    """Replace exact text inside existing OOXML text nodes, preserving package parts and styling."""
    total = 0
    output = io.BytesIO()
    with zipfile.ZipFile(path, "r") as source_zip, zipfile.ZipFile(output, "w") as target_zip:
        for info in source_zip.infolist():
            data = source_zip.read(info.filename)
            relevant = (
                path.suffix.lower() == ".pptx" and re.fullmatch(r"ppt/slides/slide\d+\.xml", info.filename)
            ) or (
                path.suffix.lower() == ".xlsx" and info.filename.startswith("xl/") and info.filename.endswith(".xml")
            )
            if relevant:
                text = data.decode("utf-8")
                # Target XML text nodes only; markup, formatting, relationships, and formulas stay byte-for-byte intact.
                pattern = re.compile(r"(<(?:[A-Za-z_][\w.-]*:)?t\b[^>]*>)(.*?)(</(?:[A-Za-z_][\w.-]*:)?t\s*>)", re.S)

                if path.suffix.lower() == ".pptx":
                    # PowerPoint commonly splits one visible phrase across
                    # several styled runs. Edit the paragraph's text nodes
                    # together while keeping all run and slide XML intact.
                    paragraph_pattern = re.compile(r"<a:p(?:\s[^>]*)?>.*?</a:p>", re.S)

                    def replace_paragraph(paragraph_match):
                        nonlocal total
                        xml = paragraph_match.group(0)
                        nodes = list(pattern.finditer(xml))
                        values = [html.unescape(node.group(2)) for node in nodes]
                        combined = "".join(values)
                        matches = list(re.finditer(re.escape(find), combined))
                        if not matches:
                            return xml
                        offsets = [0]
                        for value in values:
                            offsets.append(offsets[-1] + len(value))
                        for match in reversed(matches):
                            start, end = match.span()
                            first = next(i for i in range(len(values)) if offsets[i + 1] > start)
                            last = next(i for i in range(first, len(values)) if offsets[i + 1] >= end)
                            prefix = values[first][:start - offsets[first]]
                            suffix = values[last][end - offsets[last]:]
                            if first == last:
                                values[first] = prefix + replace + suffix
                            else:
                                values[first] = prefix + replace
                                for index in range(first + 1, last):
                                    values[index] = ""
                                values[last] = suffix
                            total += 1
                        for node, value in reversed(list(zip(nodes, values))):
                            xml = xml[:node.start(2)] + html.escape(value, quote=False) + xml[node.end(2):]
                        return xml

                    data = paragraph_pattern.sub(replace_paragraph, text).encode("utf-8")
                    target_zip.writestr(info, data)
                    continue

                def replace_node(match):
                    nonlocal total
                    visible = html.unescape(match.group(2))
                    count = visible.count(find)
                    if not count:
                        return match.group(0)
                    total += count
                    return match.group(1) + html.escape(visible.replace(find, replace), quote=False) + match.group(3)

                data = pattern.sub(replace_node, text).encode("utf-8")
            target_zip.writestr(info, data)
    if total:
        path.write_bytes(output.getvalue())
    return total


def _replace_csv_text(path: Path, find: str, replace: str) -> int:
    raw = path.read_bytes()
    has_bom = raw.startswith(b"\xef\xbb\xbf")
    source_text = raw.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(source_text[:8192])
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(source_text, newline=""), dialect))
    total = 0
    for row in rows:
        for index, value in enumerate(row):
            count = value.count(find)
            if count:
                row[index] = value.replace(find, replace)
                total += count
    if total:
        output = io.StringIO(newline="")
        writer = csv.writer(output, dialect, lineterminator="\r\n" if "\r\n" in source_text else "\n")
        writer.writerows(rows)
        path.write_bytes((b"\xef\xbb\xbf" if has_bom else b"") + output.getvalue().encode("utf-8"))
    return total


def _replace_html_text(path: Path, find: str, replace: str) -> int:
    text = path.read_text(encoding="utf-8", errors="replace")
    total = 0
    pattern = re.compile(r"(?<=>)([^<>]+)(?=<)", re.S)

    def replace_node(match):
        nonlocal total
        visible = html.unescape(match.group(1))
        count = visible.count(find)
        if not count:
            return match.group(0)
        total += count
        return html.escape(visible.replace(find, replace), quote=False)

    updated = pattern.sub(replace_node, text)
    if total:
        path.write_text(updated, encoding="utf-8", newline="")
    return total


def extract_document(path: str | Path, limit: int | None = 70000) -> tuple[str, dict]:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        import pymupdf
        with pymupdf.open(path) as doc:
            page_text, scanned_pages, page_map, headings = [], [], [], []
            images, offset = 0, 0
            for number, page in enumerate(doc, start=1):
                extracted = page.get_text("text")
                page_images = page.get_images(full=True)
                images += len(page_images)
                page_text.append(extracted)
                page_map.append({"page": number, "start": offset, "end": offset + len(extracted)})
                offset += len(extracted) + 2
                # Textless illustrated pages need OCR; short selectable headings alone
                # are not proof that their screenshot bodies have been transcribed.
                if len(extracted.strip()) < 80 and (page_images or page.get_drawings()):
                    scanned_pages.append(number)
                lines = [line.strip() for line in extracted.splitlines() if line.strip()]
                for line in lines:
                    if re.match(r"(?i)^(?:self learning|activity|chapter|section)\s*[-:–]?\s*\d+", line):
                        headings.append({"page": number, "heading": line[:180]})
            text = "\n\n".join(page_text)
            meta = {"filename": path.name, "pdf_metadata": dict(doc.metadata or {}),
                    "pages": len(doc), "page_map": page_map, "headings": headings,
                    "outline": doc.get_toc(), "images": images,
                    "scanned_pages": scanned_pages, "text_characters": len(text)}
    elif ext == ".docx":
        from docx import Document
        doc = Document(path)
        text = "\n".join(p.text for p in _docx_paragraphs(doc))
        meta = {"paragraphs": len(list(_docx_paragraphs(doc))), "images": len(doc.inline_shapes)}
    elif ext in {".pptx", ".xlsx"}:
        text, meta = _ooxml_text(path)
    elif ext == ".csv":
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            rows = list(csv.reader(handle))
        text = "\n".join("\t".join(row) for row in rows)
        meta = {"rows": len(rows), "images": 0}
    elif ext in {".html", ".htm"}:
        text = _html_text(path.read_text(encoding="utf-8", errors="replace"))
        meta = {"images": 0}
    elif ext in {".txt", ".md", ".rtf"}:
        text = path.read_text(encoding="utf-8", errors="replace")
        if ext == ".rtf":
            text = re.sub(r"\\'[0-9a-fA-F]{2}|\\[a-zA-Z]+-?\d* ?|[{}]", "", text)
        meta = {"images": 0}
    else:
        raise ValueError("Supported document formats are PDF, DOCX, PPTX, XLSX, CSV, HTML, TXT, MD, and RTF.")
    return text[:limit], meta


@lru_cache(maxsize=2)
def _extract_document_version(path: str, size: int, modified_ns: int,
                              limit: int | None) -> tuple[str, dict]:
    return extract_document(path, limit=limit)


def extract_document_cached(path: str | Path, limit: int | None = 5_000_000) -> tuple[str, dict]:
    """Reuse extracted text until the source file changes; return isolated metadata."""
    source = Path(path).resolve()
    stat = source.stat()
    before = _extract_document_version.cache_info().hits
    started = time.perf_counter()
    text, meta = _extract_document_version(str(source),stat.st_size,stat.st_mtime_ns,None)
    text=text[:limit]
    meta = copy.deepcopy(meta)
    meta['identity'] = f"{source}:{stat.st_size}:{stat.st_mtime_ns}"
    meta['parse_cache_hit'] = _extract_document_version.cache_info().hits > before
    meta['pdf_parse_ms'] = round((time.perf_counter()-started)*1000, 2)
    return text, meta


def parse_pdf_media_request(prompt: str) -> dict | None:
    """Recognize deterministic PDF-to-image requests without involving the chat model.

    Plain "show me the image from this PDF" requests render page previews. Requests
    that explicitly mention embedded images extract the original image resources.
    """
    low = prompt.lower()
    media_words = r"(?:image|picture|photo|screenshot|preview|render)"
    pdf_words = r"(?:pdf|page|document|file)"
    if not re.search(r"\b(?:show|give|send|export|render|extract|save|convert|make)\b", low):
        return None
    if not re.search(rf"\b{media_words}\b", low):
        return None
    if not re.search(rf"\b{pdf_words}\b", low) and not re.search(r"\battached\b", low):
        return None
    embedded = bool(re.search(r"\b(?:embedded|original|contained)\s+(?:image|picture|photo)s?\b|\bextract\s+(?:the\s+)?images?\b", low))
    page_match = re.search(r"\bpage\s+(\d+)\b", low)
    pages = [int(page_match.group(1))] if page_match else None
    return {"mode": "embedded" if embedded else "pages", "pages": pages}


def render_pdf_pages(path: str | Path, output_dir: str | Path, pages: list[int] | None = None,
                     max_pages: int = 12, dpi: int = 144) -> list[str]:
    """Render selected PDF pages as PNG files for inline chat preview."""
    import pymupdf
    source = Path(path)
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open(source)
    try:
        selected = pages if pages is not None else list(range(1, min(len(doc), max_pages) + 1))
        if not selected:
            raise ValueError("This PDF has no pages to render.")
        if len(selected) > max_pages:
            raise ValueError(f"To keep the export responsive, render at most {max_pages} PDF pages at a time.")
        stamp = time.strftime("%Y%m%d_%H%M%S")
        outputs = []
        for number in selected:
            if number < 1 or number > len(doc):
                raise ValueError(f"Page {number} is outside this PDF. It contains {len(doc)} page(s).")
            page = doc[number - 1]
            pix = page.get_pixmap(matrix=pymupdf.Matrix(dpi / 72, dpi / 72), alpha=False)
            target = folder / f"{source.stem}_page_{number}_{stamp}.png"
            pix.save(target)
            outputs.append(str(target))
        return outputs
    finally:
        doc.close()


def extract_pdf_images(path: str | Path, output_dir: str | Path, max_images: int = 12) -> list[str]:
    """Extract distinct embedded PDF image resources as viewable image files."""
    import pymupdf
    source = Path(path)
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    doc = pymupdf.open(source)
    try:
        refs = []
        for page in doc:
            for image in page.get_images(full=True):
                if image[0] not in refs:
                    refs.append(image[0])
        if not refs:
            raise ValueError("This PDF does not contain extractable embedded images. I can still render its pages as images.")
        if len(refs) > max_images:
            raise ValueError(f"This PDF contains {len(refs)} embedded images. Export them in groups of {max_images} or ask for a specific image number.")
        stamp = time.strftime("%Y%m%d_%H%M%S")
        outputs = []
        for index, xref in enumerate(refs, start=1):
            info = doc.extract_image(xref)
            ext = info.get("ext", "png").lower()
            if ext not in {"png", "jpg", "jpeg", "jpx", "jbig2", "bmp", "tif", "tiff", "webp"}:
                ext = "png"
            target = folder / f"{source.stem}_embedded_{index}_{stamp}.{ext}"
            target.write_bytes(info["image"])
            # Some raw PDF formats are not supported by Qt/WebEngine. Convert those to PNG.
            if ext in {"jpx", "jbig2"}:
                pix = pymupdf.Pixmap(doc, xref)
                png = target.with_suffix(".png")
                pix.save(png)
                target.unlink(missing_ok=True)
                target = png
            outputs.append(str(target))
        return outputs
    finally:
        doc.close()


def _output_copy(source: Path, output_dir: str | Path = "") -> Path:
    folder = Path(output_dir) if output_dir else source.parent / "edited"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = time.strftime("%Y%m%d_%H%M%S")
    return folder / f"{source.stem}_edited_{stamp}_{uuid.uuid4().hex[:6]}{source.suffix.lower()}"


def verify_document_output(source: str | Path, output: str | Path, operations: list[dict]) -> tuple[str, dict]:
    """Re-open a generated copy and verify requested edits before it is reported complete."""
    source_path, output_path = Path(source), Path(output)
    if source_path.resolve() == output_path.resolve():
        raise ValueError("Safety check failed: an edit must be saved to a separate file; the input was not changed.")
    if not output_path.is_file() or output_path.stat().st_size == 0:
        raise ValueError("The edited copy was not created or is empty; the original file is unchanged.")
    original_text, original_meta = extract_document(source_path)
    result_text, result_meta = extract_document(output_path)
    for operation in operations:
        kind = operation.get("type", "replace")
        if kind == "replace":
            find, replace = str(operation.get("find", "")), str(operation.get("replace", ""))
            if not find:
                continue
            normalized_original = " ".join(original_text.split())
            normalized_result = " ".join(result_text.split())
            normalized_find = " ".join(find.split())
            normalized_replace = " ".join(replace.split())
            before_count = normalized_original.count(normalized_find)
            after_count = normalized_result.count(normalized_find)
            if replace and normalized_replace not in normalized_result:
                raise ValueError(f"Verification failed: the requested replacement text was not found in the saved copy. The original is unchanged.")
            if find != replace and find not in replace and before_count and before_count <= after_count:
                raise ValueError(f"Verification failed: “{find[:80]}” is still present in the saved copy. The original is unchanged.")
        elif kind == "remove_image":
            if int(result_meta.get("images", 0)) >= int(original_meta.get("images", 0)):
                raise ValueError("Verification failed: the output does not show fewer embedded images. The original is unchanged.")
        elif kind == "replace_image":
            if int(result_meta.get("images", 0)) < 1:
                raise ValueError("Verification failed: the saved copy contains no readable embedded image. The original is unchanged.")
    return result_text, result_meta


def _save_pdf_over_open_source(doc, source: Path) -> None:
    """Serialize first, close the input handle, then safely replace the output copy."""
    try:
        data = doc.tobytes(garbage=3, deflate=True)
    finally:
        doc.close()
    source.write_bytes(data)


def _replace_docx(source: Path, find: str, replace: str) -> int:
    from docx import Document
    doc = Document(source)
    count = 0
    for paragraph in _docx_paragraphs(doc):
        runs = list(paragraph.runs)
        combined = "".join(run.text for run in runs)
        offsets = [0]
        for run in runs:
            offsets.append(offsets[-1] + len(run.text))
        matches = list(re.finditer(re.escape(find), combined))
        # Work right to left so original offsets remain valid. Leave text outside
        # the selected run range and its formatting untouched.
        for match in reversed(matches):
            start, end = match.span()
            first = next(i for i in range(len(runs)) if offsets[i + 1] > start)
            last = next(i for i in range(first, len(runs)) if offsets[i + 1] >= end)
            first_text, last_text = runs[first].text, runs[last].text
            prefix = first_text[:start - offsets[first]]
            suffix = last_text[end - offsets[last]:]
            if first == last:
                runs[first].text = prefix + replace + suffix
            else:
                runs[first].text = prefix + replace
                for index in range(first + 1, last):
                    runs[index].text = ""
                runs[last].text = suffix
            count += 1
    if count:
        doc.save(source)
    return count


def _replace_pdf(source: Path, find: str, replace: str) -> int:
    import pymupdf
    doc = pymupdf.open(source)
    hits = []
    for page in doc:
        for rect in page.search_for(find):
            best = None
            best_overlap = 0.0
            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        span_text = span.get("text", "")
                        if find not in " ".join(span_text.split()):
                            continue
                        overlap = pymupdf.Rect(span["bbox"]) & rect
                        area = max(0.0, overlap.width) * max(0.0, overlap.height)
                        if area > best_overlap:
                            best = (span, line)
                            best_overlap = area
            if best is None:
                doc.close()
                raise ValueError(
                    "This PDF splits the selected text across font runs, so its original styling cannot be identified safely. "
                    "Select a shorter exact phrase contained within one styled text run."
                )
            span, line = best
            # Keep the exact replacement capitalization requested by the user.
            replacement_text = replace
            font_name = str(span.get("font", ""))
            font_flags = int(span.get("flags", 0))
            font_size = float(span.get("size", 10))
            resource_name = ""
            normalized_font = _normalize_pdf_font_name(font_name)
            for font in page.get_fonts(full=True):
                if len(font) > 4 and _normalize_pdf_font_name(str(font[3])) == normalized_font:
                    resource_name = str(font[4]).lstrip("/")
                    break
            base_font = _pdf_base14_font(font_name, font_flags)
            insert_font = f"/{resource_name}" if resource_name else base_font
            measure_font = None
            font_file = ""
            if resource_name:
                for font in page.get_fonts(full=True):
                    if len(font) > 4 and str(font[4]).lstrip("/") == resource_name:
                        xref = int(font[0])
                        if xref > 0:
                            try:
                                extracted = doc.extract_font(xref)
                                font_data = extracted[3] if len(extracted) > 3 else b""
                                if font_data:
                                    measure_font = pymupdf.Font(fontbuffer=font_data)
                            except Exception:
                                pass
                        break
            if measure_font is not None:
                try:
                    missing_glyph = any(
                        not char.isspace() and not measure_font.has_glyph(ord(char))
                        for char in replacement_text
                    )
                except Exception:
                    missing_glyph = True
                if missing_glyph:
                    font_file = _pdf_system_font_file(font_name, font_flags)
                    if font_file:
                        import hashlib
                        insert_font = "offline" + hashlib.sha1(font_file.casefold().encode("utf-8")).hexdigest()[:10]
                        try:
                            measure_font = pymupdf.Font(fontfile=font_file)
                        except Exception:
                            font_file = ""
                    if not font_file:
                        # PDF font subsets often contain only the original name's glyphs.
                        # Use the closest built-in family/style when a replacement adds
                        # new letters, rather than rejecting an otherwise safe edit.
                        fallbacks = (base_font, "helv", "hebo", "heit", "hebi", "tiro", "tibo", "cour")
                        measure_font = None
                        for fallback in fallbacks:
                            try:
                                candidate = pymupdf.Font(fontname=fallback)
                                if all(char.isspace() or candidate.has_glyph(ord(char)) for char in replacement_text):
                                    measure_font = candidate
                                    insert_font = fallback
                                    break
                            except Exception:
                                continue
                        if measure_font is None:
                            doc.close()
                            raise ValueError("No available PDF font can display every character in the replacement text.")
            if measure_font is None:
                if resource_name:
                    # Standard PDF fonts (e.g. Helvetica-Bold) have no embedded
                    # font bytes; keep using the page's exact font resource.
                    try:
                        measure_font = pymupdf.Font(fontname=base_font)
                    except Exception:
                        measure_font = None
                else:
                    font_file = _pdf_system_font_file(font_name, font_flags)
                if measure_font is None and font_file:
                    import hashlib
                    insert_font = "offline" + hashlib.sha1(font_file.casefold().encode("utf-8")).hexdigest()[:10]
                    try:
                        measure_font = pymupdf.Font(fontfile=font_file)
                    except Exception:
                        font_file = ""
                if measure_font is None:
                    try:
                        measure_font = pymupdf.Font(fontname=base_font)
                    except Exception:
                        measure_font = pymupdf.Font(fontname="helv")
                        insert_font = "helv"
            original_width = rect.width
            replacement_width = measure_font.text_length(replacement_text, fontsize=font_size)
            if replacement_width <= 0 or original_width <= 0:
                doc.close()
                raise ValueError("The original PDF font could not measure this replacement safely.")
            # A standalone recipient (as in many certificates) can use the open
            # line to its right. Only constrain it to the original run when text
            # follows on that same line; do not squash a longer name just because
            # the old name was shorter.
            source_line_text = " ".join(str(span.get("text", "")).split())
            after_match = source_line_text.split(find, 1)[1] if find in source_line_text else ""
            right_limit = float(page.rect.x1) - 36.0
            if after_match.strip():
                right_limit = min(right_limit, float(span.get("bbox", (0, 0, right_limit, 0))[2]) - 1.0)
            source_origin_y = float(span.get("origin", (rect.x0, rect.y1))[1])
            for block in page.get_text("dict").get("blocks", []):
                for other_line in block.get("lines", []):
                    for other_span in other_line.get("spans", []):
                        if other_span is span or not str(other_span.get("text", "")).strip():
                            continue
                        other_box = pymupdf.Rect(other_span.get("bbox", (0, 0, 0, 0)))
                        other_y = float(other_span.get("origin", (other_box.x0, other_box.y1))[1])
                        same_row = abs(other_y - source_origin_y) <= max(2.0, font_size * 0.22)
                        if same_row and other_box.x0 >= rect.x1 - 1.0:
                            right_limit = min(right_limit, other_box.x0 - 2.0)
            available_width = max(0.0, right_limit - rect.x0)
            horizontal_scale = min(1.0, available_width / replacement_width)
            if horizontal_scale < 0.72:
                doc.close()
                raise ValueError(
                    "The replacement is much wider than the original text. Use a shorter replacement to preserve the document layout and font size."
                )
            color_value = int(span.get("color", 0))
            color = (
                ((color_value >> 16) & 255) / 255,
                ((color_value >> 8) & 255) / 255,
                (color_value & 255) / 255,
            )
            origin = span.get("origin", (rect.x0, rect.y1))
            direction = line.get("dir", (1, 0))
            if abs(float(direction[1])) > 0.001 or float(direction[0]) < 0.999:
                doc.close()
                raise ValueError("Style-matched text replacement currently supports horizontal PDF text only.")
            hits.append({
                "page": page.number,
                "rect": rect,
                "text": replacement_text,
                "origin": (rect.x0, float(origin[1])),
                "font": insert_font,
                "fontfile": font_file,
                "size": font_size,
                "color": color,
                "scale": horizontal_scale,
            })
    if not hits:
        doc.close()
        return 0
    for hit in hits:
        # A transparent redaction removes the original glyphs while preserving a
        # certificate's watermark, background art, and lines beneath the text.
        doc[hit["page"]].add_redact_annot(hit["rect"], fill=None, cross_out=False)
    for page in doc:
        if page.first_annot:
            page.apply_redactions(images=0, graphics=0)
    for hit in hits:
        page = doc[hit["page"]]
        morph = None
        if hit["scale"] < 0.999:
            morph = (pymupdf.Point(hit["origin"]), pymupdf.Matrix(hit["scale"], 0, 0, 1, 0, 0))
        page.insert_text(
            hit["origin"], hit["text"], fontsize=hit["size"], fontname=hit["font"],
            fontfile=hit["fontfile"] or None, color=hit["color"], morph=morph, overlay=True,
        )
    _save_pdf_over_open_source(doc, source)
    return len(hits)


def _normalize_pdf_font_name(name: str) -> str:
    """Normalize PDF subset prefixes and punctuation for span/resource matching."""
    import re
    name = re.sub(r"^[A-Z]{6}\+", "", str(name))
    return re.sub(r"[^a-z0-9]", "", name.casefold())


def _pdf_base14_font(name: str, flags: int) -> str:
    """Choose the closest built-in fallback when a PDF font resource is unavailable."""
    normalized = _normalize_pdf_font_name(name)
    sans = any(token in normalized for token in ("dejavusans", "liberationsans", "arial", "helvetica", "segoe", "calibri", "verdana", "tahoma", "sans"))
    serif = any(token in normalized for token in ("times", "serif", "cambria", "georgia", "bookman")) or (bool(flags & 4) and not sans)
    mono = any(token in normalized for token in ("courier", "mono", "consol")) or bool(flags & 8)
    bold = bool(flags & 16) or "bold" in normalized or "black" in normalized or "heavy" in normalized
    italic = bool(flags & 2) or "italic" in normalized or "oblique" in normalized
    if mono:
        return "cobi" if bold and italic else "cobo" if bold else "coit" if italic else "cour"
    if serif:
        return "tibi" if bold and italic else "tibo" if bold else "tiit" if italic else "tiro"
    return "hebi" if bold and italic else "hebo" if bold else "heit" if italic else "helv"


def _pdf_system_font_file(name: str, flags: int) -> str:
    """Find a close Windows typeface when a PDF's embedded subset lacks new glyphs."""
    import os
    normalized = _normalize_pdf_font_name(name)
    bold = bool(flags & 16) or any(token in normalized for token in ("bold", "black", "heavy"))
    italic = bool(flags & 2) or any(token in normalized for token in ("italic", "oblique"))
    style = "bold_italic" if bold and italic else "bold" if bold else "italic" if italic else "regular"
    families = (
        # DejaVu Sans is common in generated certificates and Linux-authored
        # PDFs. Its subset often lacks the replacement name's letters, so use
        # Windows Arial (a close sans-serif substitute) rather than Times.
        (("dejavusansmono", "liberationmono", "courier", "mono", "consol"), {"regular": "cour.ttf", "bold": "courbd.ttf", "italic": "couri.ttf", "bold_italic": "courbi.ttf"}),
        (("dejavusans", "liberationsans", "arial", "helvetica", "sansserif", "sans"), {"regular": "arial.ttf", "bold": "arialbd.ttf", "italic": "ariali.ttf", "bold_italic": "arialbi.ttf"}),
        (("arial", "helvetica", "liberationsans", "sansserif"), {"regular": "arial.ttf", "bold": "arialbd.ttf", "italic": "ariali.ttf", "bold_italic": "arialbi.ttf"}),
        (("times", "serif", "bookman"), {"regular": "times.ttf", "bold": "timesbd.ttf", "italic": "timesi.ttf", "bold_italic": "timesbi.ttf"}),
        (("courier", "mono", "consol"), {"regular": "cour.ttf", "bold": "courbd.ttf", "italic": "couri.ttf", "bold_italic": "courbi.ttf"}),
        (("calibri",), {"regular": "calibri.ttf", "bold": "calibrib.ttf", "italic": "calibrii.ttf", "bold_italic": "calibriz.ttf"}),
        (("cambria",), {"regular": "cambria.ttc", "bold": "cambriab.ttf", "italic": "cambriai.ttf", "bold_italic": "cambriaz.ttf"}),
        (("georgia",), {"regular": "georgia.ttf", "bold": "georgiab.ttf", "italic": "georgiai.ttf", "bold_italic": "georgiaz.ttf"}),
        (("verdana",), {"regular": "verdana.ttf", "bold": "verdanab.ttf", "italic": "verdanai.ttf", "bold_italic": "verdanaz.ttf"}),
        (("tahoma",), {"regular": "tahoma.ttf", "bold": "tahomabd.ttf", "italic": "tahoma.ttf", "bold_italic": "tahomabd.ttf"}),
    )
    windows_fonts = os.path.join(os.environ.get("WINDIR", r"C:\Windows"), "Fonts")
    for tokens, filenames in families:
        if any(token in normalized for token in tokens):
            candidate = os.path.join(windows_fonts, filenames[style])
            if os.path.isfile(candidate):
                return candidate
    return ""


def _remove_document_image(source: Path, index: int) -> int:
    if source.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(source)
        shapes = list(doc.inline_shapes)
        if index < 1 or index > len(shapes):
            raise ValueError(f"This Word document has {len(shapes)} inline images; image numbers start at 1.")
        shape = shapes[index - 1]
        shape._inline.getparent().remove(shape._inline)
        doc.save(source)
        return 1
    if source.suffix.lower() == ".pdf":
        import pymupdf
        doc = pymupdf.open(source)
        try:
            found = []
            for page in doc:
                for block in page.get_text("dict").get("blocks", []):
                    if block.get("type") == 1 and block.get("bbox"):
                        found.append((page.number, pymupdf.Rect(block["bbox"])))
            if index == 999:index=len(found)
            if index < 1 or index > len(found):
                raise ValueError(f"This PDF has {len(found)} visible images; image numbers start at 1.")
            page_number, rect = found[index - 1]
            doc[page_number].add_redact_annot(rect, fill=(1, 1, 1), cross_out=False)
            doc[page_number].apply_redactions(images=2, graphics=0)
            _save_pdf_over_open_source(doc, source)
            doc = None
            return 1
        finally:
            if doc is not None:
                doc.close()
    raise ValueError("Image removal is supported in PDF and DOCX files.")


def _replace_document_image(source: Path, index: int, replacement: Path) -> int:
    if not replacement.is_file():
        raise ValueError("The replacement image attachment could not be found.")
    if source.suffix.lower() == ".docx":
        from docx import Document
        doc = Document(source)
        shapes = list(doc.inline_shapes)
        if index < 1 or index > len(shapes):
            raise ValueError(f"This Word document has {len(shapes)} inline images; image numbers start at 1.")
        rel_id = shapes[index - 1]._inline.graphic.graphicData.pic.blipFill.blip.embed
        image_part = doc.part.related_parts[rel_id]
        image_part._blob = replacement.read_bytes()
        doc.save(source)
        return 1
    if source.suffix.lower() == ".pdf":
        import pymupdf
        doc = pymupdf.open(source)
        try:
            refs = []
            page_for_ref = {}
            for page in doc:
                for image in page.get_images(full=True):
                    if image[0] not in refs:
                        refs.append(image[0])
                        page_for_ref[image[0]] = page.number
            if index < 1 or index > len(refs):
                raise ValueError(f"This PDF has {len(refs)} embedded images; image numbers start at 1.")
            xref = refs[index - 1]
            # PyMuPDF exposes replacement on Page, not Document. Updating the
            # shared xref replaces the image wherever the PDF reuses it.
            doc[page_for_ref[xref]].replace_image(xref, filename=str(replacement))
            _save_pdf_over_open_source(doc, source)
            doc = None
            return 1
        finally:
            if doc is not None:
                doc.close()
    raise ValueError("Image replacement is supported in PDF and DOCX files.")


def apply_document_operations(path: str | Path, operations: list[dict], replacement_image: str = "", output_dir: str | Path = "") -> tuple[str, str]:
    source = Path(path)
    if not is_document(source):
        raise ValueError("This file is not a supported PDF, Word, PowerPoint, Excel, CSV, HTML, or text document.")
    out = _output_copy(source, output_dir)
    shutil.copy2(source, out)
    total = 0
    descriptions = []
    try:
        for operation in operations:
            kind = operation.get("type", "replace")
            if kind == "remove_image":
                changed = _remove_document_image(out, int(operation.get("index", 1)))
                descriptions.append(f"Removed embedded image {operation.get('index', 1)}")
            elif kind == "replace_image":
                if not replacement_image:
                    raise ValueError("Attach a second image file to replace an embedded document image.")
                changed = _replace_document_image(out, int(operation.get("index", 1)), Path(replacement_image))
                descriptions.append(f"Replaced embedded image {operation.get('index', 1)}")
            elif kind == "replace":
                find, replace = str(operation.get("find", "")), str(operation.get("replace", ""))
                if not find:
                    raise ValueError("A text replacement needs the exact text to find.")
                if out.suffix.lower() == ".docx":
                    changed = _replace_docx(out, find, replace)
                elif out.suffix.lower() == ".pdf":
                    changed = _replace_pdf(out, find, replace)
                elif out.suffix.lower() in {".pptx", ".xlsx"}:
                    changed = _replace_ooxml_text(out, find, replace)
                elif out.suffix.lower() == ".csv":
                    changed = _replace_csv_text(out, find, replace)
                elif out.suffix.lower() in {".html", ".htm"}:
                    changed = _replace_html_text(out, find, replace)
                else:
                    content = out.read_text(encoding="utf-8", errors="replace")
                    changed = content.count(find)
                    if changed:
                        out.write_text(content.replace(find, replace), encoding="utf-8")
                if not changed:
                    raise ValueError(f"Could not find this exact text in the document: {find[:160]}")
                total += changed
                descriptions.append(f"Replaced {changed} occurrence(s) of “{find[:80]}”")
            else:
                raise ValueError("Unsupported document edit operation.")
        if not descriptions:
            raise ValueError("No document changes were requested.")
        verify_document_output(source, out, operations)
    except Exception:
        try:
            out.unlink(missing_ok=True)
        except OSError:
            pass
        raise
    return str(out), "; ".join(descriptions)


def _clean_edit_value(value: str) -> str:
    value = re.sub(r"\s+", " ", str(value)).strip()
    return value.strip(" \t\r\n\"'“”‘’`").rstrip(" .?!;:").strip()


def _infer_labeled_name(document_text: str) -> str:
    """Return a name only when the document identifies it unambiguously."""
    patterns = (
        r"(?im)\b(?:full\s+)?name\s*[:：\-]\s*(?P<value>[^\r\n|;]{2,120})",
        r"(?im)^\s*(?:recipient|candidate|student|applicant|employee)\s*[:：\-]\s*(?P<value>[^\r\n|;]{2,120})",
        r"(?im)^\s*(?:full\s+name|name|recipient|candidate|student|applicant|employee)\s*[:：\-]?\s*\r?\n\s*(?P<value>[^\r\n|;]{2,120})",
        r"(?i:awarded\s+to|presented\s+to|issued\s+to|certify\s+that|dear)\s+"
        r"(?P<value>[A-Z][A-Za-z.'’\-]*(?:\s+[A-Z][A-Za-z.'’\-]*){0,3})",
    )
    values = []
    for pattern in patterns:
        for match in re.finditer(pattern, document_text or ""):
            value = _clean_edit_value(match.group("value"))
            # A label followed by a whole paragraph is not a safe name target.
            if value and 1 <= len(value.split()) <= 6:
                if value.casefold() not in {item.casefold() for item in values}:
                    values.append(value)
    if len(values) == 1:
        return values[0]
    if len(values) > 1:
        raise ValueError(
            "This document identifies more than one possible name. Quote the exact name to replace, "
            "for example: replace \"old name\" with \"new name\"."
        )
    raise ValueError(
        "I could not identify one unambiguous name field or recipient in this document. "
        "Quote the exact text to replace, for example: replace \"old name\" with \"new name\"."
    )


def _infer_labeled_field(document_text: str, field: str) -> str | None:
    """Resolve one explicit field label without asking the model to guess."""
    escaped = re.escape(field)
    candidates = []
    for match in re.finditer(rf"(?im)^\s*{escaped}\s*[:：\-]\s*(?P<value>[^\r\n|;]{{1,160}})", document_text):
        value = _clean_edit_value(match.group("value"))
        if value and value.casefold() not in {item.casefold() for item in candidates}:
            candidates.append(value)
    return candidates[0] if len(candidates) == 1 else None


def parse_direct_edit(prompt: str, replacement_image: str = "", document_text: str = "") -> list[dict] | None:
    """Parse safe, common edit requests without asking the chat model for a plan.

    This keeps exact replacements deterministic and avoids model-side structured
    output failures for simple requests such as changing a labeled PDF name.
    """
    # Explicit old/new name requests, including unquoted names.
    match = re.search(
        r"\b(?:change|replace|edit|update|set|rename)\b[^\n]{0,100}?"
        r"\b(?:full\s+)?(?:name|recipient|candidate|student|applicant|employee)\b[^\n]{0,80}?"
        r"\bfrom\s+(?P<old>.+?)\s+(?:to|with|as)\s+(?P<new>.+?)\s*[.!?]?\s*$",
        prompt, re.I,
    )
    if match:
        old, new = _clean_edit_value(match.group("old")), _clean_edit_value(match.group("new"))
        if old and new:
            return [{"type": "replace", "find": old, "replace": new}]

    quoted = r"(?:\"([^\"]+)\"|'([^']+)'|“([^”]+)”|‘([^’]+)’|`([^`]+)`)"
    match = re.search(r"\b(?:replace|change)\s+" + quoted + r"\s+(?:with|to)\s+" + quoted, prompt, re.I | re.S)
    if match:
        groups = match.groups()
        return [{"type": "replace", "find": next(x for x in groups[:5] if x), "replace": next(x for x in groups[5:] if x)}]

    # "Change the name in this PDF to …": use only an explicit, unique label
    # from selectable document text. Never guess which person the user means.
    match = re.search(
        r"\b(?:change|replace|edit|update|set|rename)\b[^\n]{0,120}?"
        r"\b(?:full\s+)?(?:name|recipient|candidate|student|applicant|employee)\b"
        r"[^\n]{0,100}?\b(?:to|with|as)\s+(?P<new>[^\n!?]+?)\s*[.!?]?\s*$",
        prompt, re.I,
    )
    if match:
        new = _clean_edit_value(match.group("new"))
        if new:
            if not str(document_text or "").strip():
                return None
            try:
                old = _infer_labeled_name(document_text)
            except ValueError:
                # Let GPT-OSS use the full document context when several fields
                # match the generic request instead of failing before planning.
                return None
            if old.casefold() == new.casefold():
                raise ValueError(f"The document already contains the requested name: {new}")
            return [{"type": "replace", "find": old, "replace": new}]

    # Handle a uniquely labeled field (date, email, company, etc.) locally.
    # Ambiguous or unlabeled targets still go to GPT-OSS with document context.
    match = re.search(
        r"\b(?:change|replace|edit|update|set)\b[^\n]{0,100}?"
        r"\b(?:the\s+)?(?P<field>date|email|phone|address|company|organization|subject)\b"
        r"[^\n]{0,100}?\b(?:to|with|as)\s+(?P<new>[^\n!?]+?)\s*[.!?]?\s*$",
        prompt, re.I,
    )
    if match and document_text.strip():
        old = _infer_labeled_field(document_text, match.group("field"))
        new = _clean_edit_value(match.group("new"))
        if old and new and old.casefold() != new.casefold():
            return [{"type": "replace", "find": old, "replace": new}]

    match = re.search(r"\bremove\s+(?:the\s+)?(?:embedded\s+)?image\s*(\d+|first|second|third|last)?\b", prompt, re.I)
    if match:
        number = {"first": 1, "second": 2, "third": 3, "last": 999}.get((match.group(1) or "1").lower(), None)
        number = int(match.group(1)) if number is None else number
        return [{"type": "remove_image", "index": number}]
    match = re.search(r"\breplace\s+(?:the\s+)?(?:embedded\s+)?image\s*(\d+|first|second|third)?\b", prompt, re.I)
    if match and replacement_image:
        token = (match.group(1) or "1").lower()
        number = {"first": 1, "second": 2, "third": 3}.get(token)
        return [{"type": "replace_image", "index": int(token) if token.isdigit() else number or 1}]
    match = re.search(r"\bdelete\s+(?:the\s+)?(?:embedded\s+)?image\s*(\d+|first|second|third)?\b", prompt, re.I)
    if match:
        token = (match.group(1) or "1").lower()
        number = {"first": 1, "second": 2, "third": 3}.get(token)
        return [{"type": "remove_image", "index": int(token) if token.isdigit() else number or 1}]

    # Exact unquoted replacement such as "change John Doe to Jane Doe".
    # Image replacement commands are handled above so they cannot be mistaken
    # for text replacement requests.
    match = re.search(
        r"\b(?:replace|change|update|edit|rename)\s+(?P<old>[^\n.!?]+?)\s+"
        r"(?:with|to|as)\s+(?P<new>[^\n.!?]+?)\s*[.!?]?\s*$",
        prompt, re.I,
    )
    if match:
        old = _clean_edit_value(match.group("old"))
        new = _clean_edit_value(match.group("new"))
        # A file-level description is not itself the old text to replace.
        generic_target = re.fullmatch(
            r"(?:(?:the|this|that|current|existing|document|pdf|file)\s+)*(?:"
            r"title|heading|header|date|name|company|organization|address|phone|email|"
            r"recipient|author|subject|body|paragraph|sentence|section|page|font|format|layout)"
            r"(?:\s+(?:field|value|text|heading|title|date|name))?",
            old, re.I,
        ) or re.search(
            r"\b(?:title|heading|header|date|name|company|organization|address|phone|email|"
            r"recipient|author|subject|body|paragraph|sentence|section|page|font|format|layout)$",
            old, re.I,
        )
        if (old and new and not generic_target and
                not re.search(r"\b(?:name|text|word|phrase|title|date|heading|header|company|address|email|phone)\b.*\b(?:pdf|document|file)\b", old, re.I)):
            return [{"type": "replace", "find": old, "replace": new}]
    return None


class OutputMode(str, Enum):
    CHAT = "CHAT"
    PDF_EXPORT = "PDF_EXPORT"
    FILE_EXPORT = "FILE_EXPORT"

    def __str__(self) -> str:
        if self.value == "CHAT":
            return "chat"
        elif self.value == "PDF_EXPORT":
            return "document"
        return "document"

    def __eq__(self, other: object) -> bool:
        if isinstance(other, str):
            if other.upper() == self.value:
                return True
            if self.value == "CHAT" and other.lower() == "chat":
                return True
            if self.value in ("PDF_EXPORT", "FILE_EXPORT") and other.lower() in ("document", "pdf_export", "file_export"):
                return True
        return super().__eq__(other)

    def __hash__(self) -> int:
        return hash(self.value)


# Negative override patterns: User explicitly specifies chat / text only or no file/pdf creation
CHAT_OVERRIDE_PATTERNS: tuple[str, ...] = (
    r"\banswer\s+(?:only\s+)?in\s+(?:the\s+)?chat(?:\s+only)?\b",
    r"\bchat\s+only\b",
    r"\bin\s+chat\s+only\b",
    r"\bin\s+(?:the\s+)?chat\b",
    r"\bno\s+pdf\b",
    r"\bdo\s+not\s+(?:make|create|generate|export|download|produce|build|save)\s+(?:a\s+|any\s+)?pdf\b",
    r"\bdon'?t\s+(?:make|create|generate|export|download|produce|build|save)\s+(?:a\s+|any\s+)?pdf\b",
    r"\bno\s+(?:file|document|attachment|artifact)\b",
    r"\bdo\s+not\s+(?:make|create|generate|export|produce|build|save)\s+(?:a\s+|any\s+)?(?:file|document)\b",
    r"\bdon'?t\s+(?:make|create|generate|export|produce|build|save)\s+(?:a\s+|any\s+)?(?:file|document)\b",
    r"\bdo\s+not\s+create\s+a\s+file\b",
    r"\bnever\s+(?:create|make|generate|export|save)\s+(?:a\s+|any\s+)?(?:pdf|file|document)\b",
    r"\btext\s+only\b",
    r"\bplain\s+text\s+only\b",
    r"\bdo\s+not\s+export\b",
    r"\bdon'?t\s+export\b",
    r"\bdisplay\s+directly(?:\s+in\s+chat)?\b",
    r"\bdisplay\s+(?:the\s+)?answer\s+in\s+chat\b",
    r"\bshow\s+(?:the\s+)?(?:answer|response|it)\s+in\s+(?:the\s+)?chat\b",
    r"\bdo\s+not\s+use\s+any\s+other\s+format\b",
    r"\bno\s+other\s+format\b",
    r"\brespond\s+in\s+chat\b",
    r"\boutput\s+in\s+chat\b",
    r"\bdeliver(?:ed)?\s+(?:only\s+)?in\s+chat\b",
    r"\bjust\s+chat\b",
    r"\bwithout\s+(?:creating|making|generating|saving)\s+(?:a\s+|any\s+)?(?:pdf|file|document)\b",
    r"\bpdf\s+created\s*:\s*no\b",
    r"\bno\s+pdf\s+created\b",
    r"\bno\s+file\s+created\b",
    r"\b(?:stay|keep)\s+in\s+chat\b",
    r"\boutput\s*mode\s*:\s*chat\b",
    r"\boutput\s*:\s*chat\b",
    r"\bformat\s*:\s*chat\b",
    r"\bchat\s+format\b",
)

# Reference or discussion of retrieved document text that should NOT trigger document creation
DOCUMENT_REFERENCE_PATTERNS: tuple[str, ...] = (
    r"\b(?:retrieved\s+document|retrieved\s+text|in\s+the\s+(?:retrieved\s+)?document|according\s+to\s+(?:the\s+)?document|document\s+(?:contains|says|mentions|states|includes))\b",
    r"\b(?:quoted|passage|excerpt|source\s+text|uploaded\s+file)\b",
)

# Coding or informational queries about PDF / DOCX / PPTX that should NOT trigger document creation
CODING_OR_INFORMATIONAL_PATTERNS: tuple[str, ...] = (
    r"\b(?:python|javascript|typescript|c\+\+|java|c#|rust|go|php|bash|powershell|script|code|program|library|module|package|api|function|class|method)\b",
    r"\b(?:how\s+(?:to|do|can)|tutorial|guide|explain|describe|what\s+is|difference\s+between|why|parse|read|extract|inspect|view|analyze|render|compress|split|merge)\b",
    r"\b(?:pypdf|pymupdf|fitz|pdfplumber|reportlab|weasyprint|pdfkit)\b",
)

# Affirmative requests to create a document file
AFFIRMATIVE_PDF_PATTERNS: tuple[str, ...] = (
    r"\b(?:create|make|generate|export|save|download|draft|build|produce)\s+(?:(?:me|us)\s+)?(?:a\s+)?(?:new\s+)?pdf\s+(?:file|document|report|presentation)\b",
    r"\b(?:create|make|generate|export|save|download|draft|build|produce)\s+(?:(?:me|us)\s+)?(?:a\s+)?(?:new\s+)?pdf\b",
    r"\b(?:export|save|download|render)(?:\s+(?:this|the\s+response|it))?\s+(?:to|as|into)\s+(?:a\s+)?pdf\b",
    r"\bpdf\s+(?:format|export|download|document)\s+(?:please|requested|only)\b",
    r"^(?:please\s+)?pdf\s*$",
)

AFFIRMATIVE_DOCX_PATTERNS: tuple[str, ...] = (
    r"\b(?:create|make|generate|export|save|download|draft|build|produce)\s+(?:(?:me|us)\s+)?(?:a\s+)?(?:new\s+)?(?:word\s+(?:document|file|doc)|docx|doc\s+file)\b",
    r"\b(?:export|save|download|render)(?:\s+(?:this|the\s+response|it))?\s+(?:to|as|into)\s+(?:a\s+)?(?:word|docx)\b",
)

AFFIRMATIVE_PPTX_PATTERNS: tuple[str, ...] = (
    r"\b(?:create|make|generate|export|save|download|draft|build|produce)\s+(?:(?:me|us)\s+)?(?:a\s+)?(?:new\s+)?(?:powerpoint|presentation|slide\s*deck|pptx?|slides)\b",
    r"\b(?:export|save|download|render)(?:\s+(?:this|the\s+response|it))?\s+(?:to|as|into)\s+(?:a\s+)?(?:powerpoint|presentation|pptx)\b",
)


def format_debug_trace(
    model_output_type: str = "text",
    validated_output_type: str = "text",
    output_mode: str = "CHAT",
    backend_response_type: str = "chat_message",
    http_content_type: str = "application/json",
    frontend_message_type: str = "text",
    artifact_created: bool = False,
    pdf_created: bool = False,
    function_that_created_pdf: str = "none",
    reason: str = "",
) -> str:
    """Format structured debug trace across the model-to-UI response pipeline."""
    return (
        f"\nMODEL OUTPUT TYPE: {model_output_type}\n"
        f"VALIDATED OUTPUT TYPE: {validated_output_type}\n"
        f"OUTPUT MODE: {output_mode}\n"
        f"BACKEND RESPONSE TYPE: {backend_response_type}\n"
        f"HTTP CONTENT TYPE: {http_content_type}\n"
        f"FRONTEND MESSAGE TYPE: {frontend_message_type}\n"
        f"ARTIFACT CREATED: {str(artifact_created).lower()}\n"
        f"PDF CREATED: {str(pdf_created).lower()}\n"
        f"FUNCTION THAT CREATED PDF: {function_that_created_pdf}\n"
        f"REASON: {reason}"
    )


def resolve_output_mode(prompt: str, has_attachment: bool = False) -> tuple[OutputMode, str, str]:
    """Resolve whether the response should be output directly in chat or saved as a document.

    Returns (output_mode, format_extension, reason) where output_mode is an OutputMode enum.
    Default is ALWAYS OutputMode.CHAT. Negative user instructions always override any file triggers.
    """
    if has_attachment:
        return OutputMode.CHAT, '', 'attachment_present'
    low = (prompt or '').strip().lower()
    if not low:
        return OutputMode.CHAT, '', 'empty_prompt'

    # Check negative overrides first: explicit negative instruction always wins
    for pat in CHAT_OVERRIDE_PATTERNS:
        if re.search(pat, low, re.I):
            return OutputMode.CHAT, '', f'negative_override: {pat}'

    # Check if prompt is discussing or querying a retrieved document
    is_doc_reference = any(re.search(pat, low, re.I) for pat in DOCUMENT_REFERENCE_PATTERNS)
    is_coding_info = any(re.search(pat, low, re.I) for pat in CODING_OR_INFORMATIONAL_PATTERNS)

    # Strip quoted substrings to prevent quoted mentions (e.g. 'create PDF') from triggering file creation
    unquoted = re.sub(r"['\"][^'\"]*['\"]", " ", low)

    formats = (
        ('.pdf', AFFIRMATIVE_PDF_PATTERNS, 'affirmative_pdf', OutputMode.PDF_EXPORT),
        ('.docx', AFFIRMATIVE_DOCX_PATTERNS, 'affirmative_docx', OutputMode.FILE_EXPORT),
        ('.pptx', AFFIRMATIVE_PPTX_PATTERNS, 'affirmative_pptx', OutputMode.FILE_EXPORT),
    )

    for ext, patterns, reason_prefix, mode_enum in formats:
        for pat in patterns:
            # Must match in unquoted text so quoted phrases from documents do not trigger export
            m = re.search(pat, unquoted, re.I)
            if m:
                if is_coding_info or is_doc_reference:
                    # If prompt is asking about coding/parsing/info or retrieved document, only trigger if command begins affirmatively
                    if not re.match(
                        r'^\s*(?:please\s+)?(?:create|make|generate|export|save|download|draft|produce)\s+(?:(?:me|us)\s+)?(?:a\s+)?(?:new\s+)?(?:pdf|word|docx|powerpoint|pptx|presentation|slide\s*deck)\b',
                        low,
                        re.I
                    ):
                        continue
                return mode_enum, ext, f'{reason_prefix}: {pat}'

    return OutputMode.CHAT, '', 'default_chat'


def requested_document_format(prompt: str, has_attachment: bool = False) -> str:
    """Return a requested output format only for an explicit file-creation task."""
    mode, ext, _ = resolve_output_mode(prompt, has_attachment)
    return ext if mode in (OutputMode.PDF_EXPORT, OutputMode.FILE_EXPORT) else ""


def assistant_document_attachment_allowed(prompt: str, source_attachment: str = "") -> bool:
    """Only display a saved document for an explicit export or edit request.

    This also hides stale PDF paths written by older releases after a chat answer.
    The saved file is left untouched; the decision affects presentation only.
    """
    mode, extension, _ = resolve_output_mode(prompt)
    if mode != OutputMode.CHAT and extension:
        return True
    if not source_attachment or not is_document(source_attachment):
        return False
    edit_intent = bool(re.search(
        r"\b(edit|rewrite|improve|correct|polish|fix|change|replace|remove|delete|update|rename|redact|shorten|expand|rephrase|convert)\b",
        prompt or "", re.I,
    ) or re.search(r"\b(make|turn)\s+(?:the|this|it|my|a)\b", prompt or "", re.I))
    return edit_intent and not any(re.search(pattern, prompt or "", re.I) for pattern in CHAT_OVERRIDE_PATTERNS)


def _document_blocks(markdown: str) -> list[tuple[str, str]]:
    """Parse model-authored Markdown into a small, predictable layout vocabulary."""
    result = []
    in_code = False
    for raw in markdown.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if line.startswith("```"):
            in_code = not in_code
            continue
        if not line:
            if result and result[-1] != ("space", ""):
                result.append(("space", ""))
            continue
        if re.fullmatch(r"[-_*]{3,}", line):
            result.append(("space", ""))
            continue
        match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if match:
            result.append(("title" if len(match.group(1)) == 1 else "heading", match.group(2).strip()))
            continue
        match = re.match(r"^(?:[-*•]|\d+[.)])\s+(.+)$", line)
        if match:
            result.append(("bullet", match.group(1).strip()))
        else:
            result.append(("body", line if not in_code else raw.rstrip()))
    return result


def _created_file_path(title: str, extension: str, output_dir: str | Path) -> Path:
    folder = Path(output_dir)
    folder.mkdir(parents=True, exist_ok=True)
    stem = re.sub(r"[^\w-]+", "_", title, flags=re.UNICODE).strip("_")[:60] or "document"
    return folder / f"{stem}_{time.strftime('%Y%m%d_%H%M%S')}_{uuid.uuid4().hex[:6]}{extension}"


def literal_document_request(prompt: str, extension: str) -> dict | None:
    """Recognize explicit literal-body exports, not open-ended writing requests.

    Everything after the marked colon is user data. Never ask a language model
    to expand text the user explicitly supplied verbatim. Ambiguous requests
    retain the normal authoring path; presentations require their slide layout.
    """
    if extension not in {'.pdf', '.docx'} or resolve_output_mode(prompt)[0] == OutputMode.CHAT:
        return None
    match = re.search(
        r'\b(?:exactly|verbatim)\s+(?:this|the following)\s+(?:sentence|text|content)'
        r'(?:\s+(?:in|as)\s+(?:its|the|document\'s)\s+body)?\s*:\s*(.+)\Z',
        prompt, re.I | re.S,
    )
    if not match or not match[1].strip() or len(match[1]) > 100_000:
        return None
    prefix = prompt[:match.start()]
    title = re.search(r'\btitled\s+(.+?)\s+(?:with|and|containing)\s*$', prefix, re.I)
    # Without an unambiguous title/body boundary, let ordinary authoring handle
    # the request rather than silently guessing where the user's title ends.
    if not title or len(title[1]) > 160 or '\n' in title[1]:
        return None
    return {'title':title[1].strip().strip('"“”'), 'body':match[1].strip()}


def create_document_from_markdown(markdown: str, extension: str, output_dir: str | Path, output_mode: OutputMode | str = OutputMode.PDF_EXPORT, *, literal_content: dict | None = None) -> tuple[str, str]:
    """Create and re-read a local PDF, DOCX, or PPTX from GPT-authored text."""
    if output_mode == OutputMode.CHAT or str(output_mode).upper() == "CHAT":
        pdf_created = False
        assert not pdf_created, "Hard assertion failed: PDF creation attempted while output_mode is CHAT"
        raise ValueError("Cannot create document: output_mode is CHAT")

    if extension not in {".pdf", ".docx", ".pptx"}:
        raise ValueError("Create a PDF, Word document, or PowerPoint presentation.")
    if literal_content is not None:
        if extension not in {'.pdf', '.docx'}:
            raise ValueError('Literal-body output supports PDF and Word documents.')
        if set(literal_content) != {'title','body'} or not all(isinstance(v,str) and v.strip() for v in literal_content.values()):
            raise ValueError('Literal document content requires a nonempty title and body.')
        blocks = [('title',literal_content['title']),('body',literal_content['body'])]
    else:
        blocks = _document_blocks(markdown)
    meaningful = [(kind, value) for kind, value in blocks if kind != "space" and value]
    if not meaningful or (literal_content is None and sum(len(value) for _, value in meaningful) < 40):
        raise ValueError("The local chat model did not produce enough document content to save safely.")
    title = next((value for kind, value in meaningful if kind == "title"), meaningful[0][1])
    target = _created_file_path(title, extension, output_dir)
    try:
        if extension == ".docx":
            _create_docx(target, blocks)
        elif extension == ".pptx":
            _create_pptx(target, blocks)
        else:
            _create_pdf(target, blocks)
        extracted, meta = extract_document(target)
        if not target.is_file() or target.stat().st_size < 500 or (literal_content is None and len(extracted.strip()) < 35):
            raise ValueError("The generated document failed its saved-file verification.")
        if literal_content is not None:
            expected = ' '.join((literal_content['title']+' '+literal_content['body']).split())
            if ' '.join(extracted.split()) != expected:
                raise ValueError('The saved document did not preserve the exact requested title and body.')
        if " ".join(title[:40].split()).casefold() not in " ".join(extracted.split()).casefold():
            raise ValueError("The saved document does not contain its requested title.")
        amount = meta.get("slides", meta.get("pages", 1))
        unit = "slide" if extension == ".pptx" else "page"
        return str(target), f"Created {target.name} ({amount} {unit}{'' if amount == 1 else 's'})."
    except Exception:
        target.unlink(missing_ok=True)
        raise


def _create_docx(target: Path, blocks: list[tuple[str, str]]) -> None:
    from docx import Document
    from docx.shared import Inches, Pt
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.9)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(11)
    for kind, value in blocks:
        if kind == "space":
            continue
        if kind == "title":
            document.add_heading(value, level=0 if not document.paragraphs else 1)
        elif kind == "heading":
            document.add_heading(value, level=2)
        elif kind == "bullet":
            document.add_paragraph(value, style="List Bullet")
        else:
            document.add_paragraph(value)
    document.save(target)


def _create_pdf(target: Path, blocks: list[tuple[str, str]]) -> None:
    import pymupdf
    doc = pymupdf.open()
    page = None
    y = 0.0
    page_width, page_height, margin = 595.0, 842.0, 52.0
    font_path = Path(r"C:\Windows\Fonts\arial.ttf")
    bold_path = Path(r"C:\Windows\Fonts\arialbd.ttf")
    font_regular = pymupdf.Font(fontfile=str(font_path)) if font_path.is_file() else pymupdf.Font(fontname="helv")
    font_bold = pymupdf.Font(fontfile=str(bold_path)) if bold_path.is_file() else pymupdf.Font(fontname="hebo")

    def new_page():
        nonlocal page, y
        page = doc.new_page(width=page_width, height=page_height)
        if font_path.is_file():
            page.insert_font(fontname="DocRegular", fontfile=str(font_path))
        if bold_path.is_file():
            page.insert_font(fontname="DocBold", fontfile=str(bold_path))
        y = margin

    def wrap(value: str, font, size: float, width: float) -> list[str]:
        words = value.split()
        lines, line = [], ""
        for word in words:
            candidate = f"{line} {word}" if line else word
            if line and font.text_length(candidate, fontsize=size) > width:
                lines.append(line)
                line = word
            else:
                line = candidate
        if line:
            lines.append(line)
        return lines or [""]

    new_page()
    for kind, value in blocks:
        if kind == "space":
            y += 7
            continue
        size = 21 if kind == "title" else 15 if kind == "heading" else 11
        advance = size * 1.38
        font = font_bold if kind in {"title", "heading"} else font_regular
        font_name = ("DocBold" if bold_path.is_file() else "hebo") if kind in {"title", "heading"} else ("DocRegular" if font_path.is_file() else "helv")
        indent = 15 if kind == "bullet" else 0
        lines = wrap(value, font, size, page_width - 2 * margin - indent)
        for number, line in enumerate(lines):
            if y + advance > page_height - margin:
                new_page()
            if kind == "bullet" and number == 0:
                page.draw_circle(pymupdf.Point(margin + 4, y - 4), 2.1, color=(0.18, 0.4, 0.75), fill=(0.18, 0.4, 0.75))
            page.insert_text(pymupdf.Point(margin + indent, y), line, fontsize=size, fontname=font_name,
                             color=(0.09, 0.13, 0.2) if kind in {"title", "heading"} else (0.16, 0.19, 0.24))
            y += advance
        if kind in {"title", "heading"}:
            y += 6
    doc.save(target, garbage=3, deflate=True)
    doc.close()


def _create_pptx(target: Path, blocks: list[tuple[str, str]]) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    slides = []
    title, body = "", []
    deck_title=""
    explicit_subslides=any(kind=='heading' and re.match(r'^slide\s+\d+\b',value,re.I) for kind,value in blocks)
    for kind, value in blocks:
        slide_heading=kind=='title' or (kind=='heading' and explicit_subslides and re.match(r'^slide\s+\d+\b',value,re.I))
        if slide_heading and title:
            # A deck-level title followed by explicit numbered slide headings
            # is metadata, not an extra empty slide. Preserve meaningful intro.
            if body or not explicit_subslides:slides.append((title, body))
            elif not deck_title:deck_title=title
            title, body = value, []
        elif slide_heading:
            title = value
        elif kind != "space":
            body.append(value)
    if title or body:
        slides.append((title or "Presentation", body))
    # A model may supply one heading and a long outline. Paginate it rather
    # than allowing the slide text to run beyond the canvas.
    paginated = []
    for slide_title, items in slides:
        for start in range(0, max(1, len(items)), 6):
            paginated.append((slide_title if start == 0 else f"{slide_title} (cont.)", items[start:start + 6]))
    presentation.core_properties.title=deck_title or (paginated[0][0] if paginated else 'Presentation')
    for slide_index,(slide_title, items) in enumerate(paginated[:40]):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(10, 17, 29)
        heading = slide.shapes.add_textbox(Inches(0.8), Inches(0.62), Inches(11.8), Inches(1.0)).text_frame
        heading.text = f'{deck_title} — {slide_title}' if slide_index==0 and deck_title else slide_title
        heading.paragraphs[0].font.size = Pt(30)
        heading.paragraphs[0].font.bold = True
        heading.paragraphs[0].font.color.rgb = RGBColor(245, 248, 252)
        body_frame = slide.shapes.add_textbox(Inches(1.05), Inches(2.0), Inches(11.15), Inches(4.8)).text_frame
        body_frame.word_wrap = True
        for index, item in enumerate(items):
            paragraph = body_frame.paragraphs[0] if index == 0 else body_frame.add_paragraph()
            paragraph.text = item
            paragraph.font.size = Pt(22 if len(item) < 150 else 18)
            paragraph.font.color.rgb = RGBColor(222, 232, 246)
            paragraph.space_after = Pt(16)
    presentation.save(target)


def export_conversations(records: list[dict], destination: str | Path, include_attachments=False):
    """Explicit local export. Write atomically; attachment originals are never altered."""
    import json
    import os
    import tempfile
    target=Path(destination)
    extension=target.suffix.lower()
    if extension not in {'.md','.txt','.json','.pdf','.zip'}:raise ValueError('Choose Markdown, TXT, JSON, PDF or ZIP.')
    def markdown(chat):
        rows=[f"# {chat['title']}", 'Exported locally from Offline GPT.']
        for message in chat['messages']:
            rows += ['\n## '+('User' if message['role']=='user' else 'Offline GPT'),message['content']]
            if message.get('attachments'):
                rows.append('Attachments: '+', '.join(Path(p).name for p in message['attachments']))
        return '\n\n'.join(rows)
    with tempfile.TemporaryDirectory(prefix='offline-gpt-export-',dir=target.parent) as folder:
        temp=Path(folder)/target.name
        if extension=='.zip':
            with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as archive:
                archive.writestr('conversations.json',json.dumps(records,ensure_ascii=False,indent=2))
                for index,chat in enumerate(records,1):
                    archive.writestr(f"{index:03d}-conversation.md",markdown(chat))
                    if include_attachments:
                        seen=set()
                        for message in chat['messages']:
                            for raw in message.get('attachments',[]):
                                source=Path(raw)
                                if source.is_file() and str(source.resolve()) not in seen:
                                    seen.add(str(source.resolve()))
                                    archive.write(source,f"attachments/{index:03d}/{len(seen):03d}-{source.name}")
        elif extension=='.json':temp.write_text(json.dumps(records,ensure_ascii=False,indent=2),encoding='utf-8')
        elif extension=='.pdf':
            created,_=create_document_from_markdown('\n\n---\n\n'.join(markdown(c) for c in records),'.pdf',folder)
            Path(created).replace(temp)
        else:
            body='\n\n---\n\n'.join(markdown(c) for c in records)
            if extension=='.txt':body=re.sub(r'(?m)^#{1,6}\s*','',body)
            temp.write_text(body,encoding='utf-8')
        os.replace(temp,target)
    return str(target)
