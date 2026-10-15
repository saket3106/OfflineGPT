"""Regressions discovered by genuine October 15 installed inference."""
import unittest
import tempfile
import threading
import json
import sqlite3
from pathlib import Path
from unittest.mock import Mock,patch
from grounding import decompose_request_tasks,is_routing_classification
from router import route_request
from intelligence import analyze_request,resolve_deterministic_task
from integrated_runtime import _local_memory_profile

class InstalledInferenceRegressions(unittest.TestCase):
    def test_historical_change_coverage_is_not_operation_completion(self):
        from conversation_memory import history_change_evidence_missing
        sources=[{'id':'M77','evidenceKind':'APPLICATION_IMAGE_TELEMETRY','excerpt':'An image completed.','runtimeFacts':{'duration_seconds':47.3}}]
        self.assertTrue(history_change_evidence_missing('Which improvements actually worked?',sources))
        self.assertFalse(history_change_evidence_missing('Which changes did I request?',sources))
        self.assertFalse(history_change_evidence_missing('How long did that image take?',sources))
        self.assertTrue(history_change_evidence_missing('Which improvements actually worked?',sources+[{'id':'M79','evidenceKind':'USER_NOTE','excerpt':'The parser has not been fixed.'}]))
        sources.append({'id':'M78','evidenceKind':'USER_STATEMENT','excerpt':'I tested the updated parser; the syntax error is fixed.'})
        self.assertFalse(history_change_evidence_missing('Which improvements actually worked?',sources))
    def test_history_finalization_cannot_invent_milestones_from_image_records(self):
        from app import ChatWorker
        from intelligence import analyze_request
        prompt='Summarize the project improvements that actually worked.'
        source={'id':'M77','chatId':'fixture','url':'offlinegpt://chat/fixture?message=77','date':'2026-10-01','created':1,'evidenceKind':'APPLICATION_IMAGE_TELEMETRY','excerpt':'An image completed.','runtimeFacts':{'duration_seconds':47.3}}
        class Engine:
            cfg={}
            def ensure_chat(self,*args,**kwargs):return 'local'
            def chat_completion(self,*args,**kwargs):return 'We now have instant response caching [M77].'
        analysis=analyze_request(prompt);analysis.update(memory_request=True,memory_sources=[source],max_retries=0)
        worker=ChatWorker('local',[{'role':'user','content':prompt}],Engine(),analysis)
        final=[];failed=[];metrics=[];worker.done.connect(final.append);worker.failed.connect(failed.append);worker.metrics.connect(metrics.append);worker.run()
        self.assertFalse(failed);self.assertEqual(len(final),1);self.assertIn('partial evidence view',final[0]);self.assertNotIn('instant response caching',final[0]);self.assertIn('M77',final[0]);self.assertEqual(metrics[-1]['validation_status'],'partial')
    def test_numbered_markdown_slide_headings_remain_separate_slides(self):
        from document_tools import create_document_from_markdown
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as directory:
            markdown='# Deck overview\n\n## Slide 1: Alpha\n- First detail\n- Second detail\n\n## Slide 2: Beta\n- Third detail\n\n## Slide 3: Gamma\n- Fourth detail'
            path,_=create_document_from_markdown(markdown,'.pptx',directory)
            deck=Presentation(path);self.assertEqual(len(deck.slides),3)
            texts=[' '.join(s.text for s in slide.shapes if s.has_text_frame) for slide in deck.slides]
            self.assertIn('Alpha',texts[0]);self.assertIn('Beta',texts[1]);self.assertIn('Gamma',texts[2])
            self.assertIn('Deck overview',texts[0])
            path,_=create_document_from_markdown('# Alpha\n- First detail for this slide\n# Beta\n- Second detail for this slide','.pptx',directory)
            self.assertEqual(len(Presentation(path).slides),2)
    def test_negated_image_operations_do_not_override_generative_edit(self):
        from image_pipeline import select_image_operation,ImageOperation
        for prohibition in ['do not upscale','never crop','avoid denoise',"don't resize"]:
            action,_=select_image_operation('Use Generative AI Edit. Replace the tabletop with wood; '+prohibition+'.',True)
            self.assertEqual(action,ImageOperation.GENERATIVE_IMAGE_EDIT)
        self.assertEqual(select_image_operation('Upscale the image without changing the subject.',True)[0],ImageOperation.UPSCALE)
    def test_domain_classification_not_request_routing(self):
        prompt='1. Classify each item as FRUIT, ANIMAL or TOOL:\nA. apple\nB. cat\nC. hammer\nD. pear\nE. dog\nF. wrench'
        analysis=analyze_request(prompt)
        self.assertEqual(analysis['expected_task_ids'],['1.'+x for x in 'ABCDEF'])
        for task in analysis['task_analyses']:
            self.assertNotEqual(task['route_target'],'CLASSIFICATION_ONLY')
            self.assertIsNone(resolve_deterministic_task(task))
    def test_domain_classify_only_not_meta_routing(self):
        for text in ('Classify only as RED or BLUE: this object.','Classify only as FRUIT, ANIMAL or TOOL: apple. Do not execute anything.','Classify each item as FRUIT, ANIMAL or TOOL. METHOD: DIRECT.'):
            self.assertFalse(is_routing_classification(text))
            self.assertNotEqual(route_request(text)['route_target'],'CLASSIFICATION_ONLY')
    def test_explicit_request_routing_preserved(self):
        self.assertTrue(is_routing_classification('Classify each request as DIRECT or TOOL:'))
        self.assertTrue(is_routing_classification('Classify as DIRECT or TOOL based on the required operation.'))
        self.assertTrue(is_routing_classification('Classify as DIRECT or TOOL, then return the labels.'))
        self.assertTrue(is_routing_classification('Classify only: Calculate 23 * 41. Do not execute it.'))
    def test_unseen_domain_categories_preserve_all_ids(self):
        tasks=decompose_request_tasks('1. Classify each sample into ACID or BASE:\nA. sample alpha\nB. sample beta\nC. sample gamma')
        self.assertEqual([t['id'] for t in tasks],['1.A','1.B','1.C'])
        self.assertTrue(all(not t['classification_only'] for t in tasks))
    def test_small_machine_bounds_context_and_cpu_profile(self):
        gib=1024**3
        p=_local_memory_profile({'total':16*gib,'available':8*gib,'pagefile_available':20*gib},4096)
        self.assertEqual(p['context'],2048);self.assertLessEqual(p['batch'],64);self.assertLessEqual(p['threads'],4)
        self.assertEqual(p['load_mode'],'none')
    def test_larger_machine_keeps_requested_context(self):
        gib=1024**3
        p=_local_memory_profile({'total':32*gib,'available':20*gib,'pagefile_available':30*gib},4096)
        self.assertEqual(p['context'],4096)
        self.assertEqual(p['load_mode'],'mmap')
    def test_document_exports_respect_isolated_data_profile(self):
        from pathlib import Path
        source=(Path(__file__).resolve().parents[1]/'app.py').read_text(encoding='utf-8')
        self.assertNotIn("APP_DIR/'generated'/'documents'",source)
        self.assertNotIn("APP_DIR / 'generated' / 'documents'",source)
    def test_native_launch_uses_measured_thread_and_load_profile(self):
        from integrated_runtime import _LocalLlamaServer
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);model=root/'fixture.gguf';model.write_bytes(b'fixture')
            runtime=_LocalLlamaServer(root);process=Mock();process.poll.return_value=None
            response=Mock();response.__enter__=Mock(return_value=response);response.__exit__=Mock(return_value=False);response.status=200
            with patch.object(runtime,'stop'),patch.object(runtime,'_server_binary',return_value=root/'llama-server.exe'),patch.object(runtime,'_check_runtime_files'),patch('integrated_runtime.subprocess.Popen',return_value=process) as launch,patch('integrated_runtime.urllib.request.urlopen',return_value=response),patch('integrated_runtime.physical_memory_status',return_value={'available':4*1024**3}):
                try:runtime.start(model,2048,'auto',lambda _:None,memory_profile={'batch':64,'ubatch':32,'threads':4,'load_mode':'none'})
                finally:
                    if runtime.log_handle:runtime.log_handle.close()
            args=launch.call_args.args[0]
            self.assertEqual(args[args.index('--load-mode')+1],'none')
            self.assertEqual(args[args.index('-t')+1],'4')
            self.assertEqual(args[args.index('-tb')+1],'4')
            self.assertEqual(args[args.index('--cache-ram')+1],'128')
            self.assertNotIn('--cache-reuse',args)
    def test_retry_remains_self_contained_after_history_trim(self):
        from intelligence import build_correction_messages
        from integrated_runtime import _LocalLlamaServer
        original='Explain safe concurrent publication of two related local files, including failure recovery.'
        messages=[{'role':'system','content':'Answer the task.'},{'role':'user','content':original}]
        retry=build_correction_messages(messages,'old draft '*1500,['Expected at most 250 words.'])
        runtime=_LocalLlamaServer(Path.cwd());runtime.context=2048
        with patch.object(runtime,'_prompt_tokens',side_effect=lambda msgs:3000 if len(msgs)>2 else 400):
            fitted,_,_,shortened=runtime._fit_messages(retry,800,lambda _:None,None)
        self.assertTrue(shortened)
        self.assertIn(original,fitted[-1]['content'])
        self.assertIn('Expected at most 250 words.',fitted[-1]['content'])
        self.assertNotIn('old draft',fitted[-1]['content'])
    def test_prompt_echo_cannot_finalize_as_empty_success(self):
        from app import ChatWorker
        from intelligence import analyze_request
        prompt='Explain how an iterator differs from a list.'
        class Engine:
            cfg={}
            def __init__(self):self.calls=0
            def ensure_chat(self,*args,**kwargs):return 'local'
            def chat_completion(self,messages,**kwargs):
                self.calls+=1
                return prompt if self.calls==1 else 'An iterator yields items on demand; a list stores its elements.'
        engine=Engine();worker=ChatWorker('local',[{'role':'user','content':prompt}],engine,analyze_request(prompt))
        final=[];failed=[];worker.done.connect(final.append);worker.failed.connect(failed.append);worker.run()
        self.assertEqual(failed,[]);self.assertEqual(engine.calls,2)
        self.assertEqual(final,['An iterator yields items on demand; a list stores its elements.'])
    def test_persistent_echo_is_explicit_failure_never_passed(self):
        from app import ChatWorker
        from intelligence import analyze_request
        prompt='Explain how an iterator differs from a list.'
        class Engine:
            cfg={}
            def ensure_chat(self,*args,**kwargs):return 'local'
            def chat_completion(self,*args,**kwargs):return prompt
        engine=Engine()
        worker=ChatWorker('local',[{'role':'user','content':prompt}],engine,analyze_request(prompt))
        final=[];failed=[];metrics=[];worker.done.connect(final.append);worker.failed.connect(failed.append);worker.metrics.connect(metrics.append);worker.run()
        self.assertFalse(final);self.assertEqual(len(failed),1);self.assertIn('no displayable answer',failed[0])
        self.assertFalse(any(m.get('validation_status')=='passed' for m in metrics))
    def test_document_reasoning_disables_agent_tools(self):
        from workspace_services import answer_documents
        workspace=Mock();workspace.search.return_value={'status':'FOUND','matches':[{'chunkId':'one','filename':'manual.pdf','page':1,'section':None,'snippet':'The value is 37.'}]}
        engine=Mock();engine.cfg={'chat_context':2048};engine.chat_completion.return_value='The value is 37 [1].'
        result=answer_documents(workspace,engine,{'query':'What is the value?'},threading.Event())
        self.assertEqual(result['answer'],'The value is 37 [1].')
        self.assertEqual(engine.chat_completion.call_args.kwargs['allowed_tools'],[])
        engine.chat_completion.return_value=object()
        with self.assertRaisesRegex(RuntimeError,'non-text'):answer_documents(workspace,engine,{'query':'What is the value?'},threading.Event())
    def test_workspace_serialization_failure_is_terminal_event(self):
        from app import WorkspaceTask
        workspace=Mock();workspace.documents.return_value={'bad':object()}
        worker=WorkspaceTask('serialize','documents',{},workspace,Mock());events=[];worker.done.connect(events.append);worker.run()
        self.assertEqual(len(events),1);result=json.loads(events[0]);self.assertFalse(result['ok']);self.assertIn('serializ',result['error'])
    def test_worker_fatal_hook_never_opens_widget(self):
        import app
        with tempfile.TemporaryDirectory() as directory,patch.object(app,'DATA_DIR',Path(directory)),patch.object(app,'_fatal_error_relay') as relay,patch.object(app.QApplication,'instance') as instance,patch.object(app.QThread,'currentThread',return_value=object()),patch.object(app.QMessageBox,'critical') as dialog:
            instance.return_value.thread.return_value=object()
            app.fatal_handler(ValueError,ValueError('test failure'),None)
            dialog.assert_not_called();relay.message.emit.assert_called_once()
            self.assertIn('test failure',(Path(directory)/'offline_gpt_crash.log').read_text())
    def test_export_budget_larger_than_context_preserves_prompt(self):
        from integrated_runtime import _LocalLlamaServer
        runtime=_LocalLlamaServer(Path.cwd());runtime.context=2048
        messages=[{'role':'system','content':'Preserve all document requirements. '*30},{'role':'user','content':'Create a short checklist.'}]
        with patch.object(runtime,'_prompt_tokens',return_value=143):
            fitted,count,output,shortened=runtime._fit_messages(messages,2400,lambda _:None,None)
        self.assertEqual(fitted,messages);self.assertEqual(count,143);self.assertEqual(output,2048-143-16);self.assertFalse(shortened)
    def test_oversized_current_request_still_fails_safely(self):
        from integrated_runtime import _LocalLlamaServer
        runtime=_LocalLlamaServer(Path.cwd());runtime.context=2048
        with patch.object(runtime,'_prompt_tokens',return_value=3000):
            with self.assertRaisesRegex(RuntimeError,'Shorten'):runtime._fit_messages([{'role':'user','content':'A long indivisible request'}],2400,lambda _:None,None)
    def test_retrieval_does_not_drop_other_pages_or_disjoint_spans(self):
        from workspace_services import LocalWorkspace
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'source.txt';path.write_text('test source')
            con=sqlite3.connect(':memory:')
            try:
                con.execute('CREATE TABLE workspace_documents(id TEXT,path TEXT)')
                con.execute('CREATE TABLE document_chunks(id TEXT,document_id TEXT,page INTEGER,section TEXT,start INTEGER,end INTEGER,text TEXT)')
                con.execute('INSERT INTO workspace_documents VALUES(?,?)',('doc',str(path)))
                for cid,page,start,text in [('a',1,0,'Inspection value 73 days.'),('b',2,30,'Inspection value 91 newton metres.'),('c',2,100,'Inspection value 53 degrees.')]:
                    con.execute('INSERT INTO document_chunks VALUES(?,?,?,?,?,?,?)',(cid,'doc',page,None,start,start+len(text),text))
                ws=object.__new__(LocalWorkspace);ws.con=con;ws.lock=threading.RLock()
                matches=ws.search('inspection value')['matches']
                self.assertEqual({m['chunkId'] for m in matches},{'a','b','c'})
            finally:con.close()
    def test_history_duration_must_match_cited_record(self):
        from conversation_memory import history_provenance_failures
        sources=[{'id':'M53','evidenceKind':'APPLICATION_IMAGE_TELEMETRY','runtimeFacts':{'duration_seconds':91.25}}, {'id':'M61','evidenceKind':'APPLICATION_IMAGE_TELEMETRY','runtimeFacts':{'duration_seconds':112.4}}]
        self.assertTrue(history_provenance_failures('[M53] Generation duration: 112 seconds.',sources))
        self.assertFalse(history_provenance_failures('[M53] Generation duration: 91 seconds.',sources))
        self.assertFalse(history_provenance_failures('[M53] Generation duration: 91250 ms.',sources))
        self.assertFalse(history_provenance_failures('[M53] Generation did not take 112 seconds.',sources))
        self.assertFalse(history_provenance_failures('[M53] took 91.25 seconds; [M61] took 112.4 seconds.',sources))
    def test_history_error_report_cannot_be_denied_as_no_problem(self):
        from conversation_memory import history_provenance_failures
        sources=[{'id':'M53','evidenceKind':'USER_ERROR_REPORT','excerpt':'Document error: replacement does not fit.'}]
        self.assertTrue(history_provenance_failures('None of these messages mention problems. [M53]',sources))
        self.assertFalse(history_provenance_failures('The error report [M53] does not establish that its fix was implemented.',sources))
    def test_history_ordered_multiline_durations_cannot_borrow_other_sources(self):
        from conversation_memory import history_provenance_failures
        sources=[{'id':i,'evidenceKind':'APPLICATION_IMAGE_TELEMETRY','runtimeFacts':{'duration_seconds':value}} for i,value in [('M21',47.3),('M22',58.6),('M23',69.7)]]
        self.assertTrue(history_provenance_failures('- Runs [M21] and [M22].\n  - Durations: 47.3 s and 69.7 s respectively.',sources))
        self.assertTrue(history_provenance_failures('- Runs [M21] and [M22].\n  - Durations: 58.6 s and 47.3 s respectively.',sources))
        self.assertFalse(history_provenance_failures('- Runs [M21] and [M22].\n  - Durations: 47.3 s and 58.6 s respectively.',sources))
    def test_history_operation_records_do_not_prove_software_improvements(self):
        from conversation_memory import history_provenance_failures
        sources=[{'id':'M21','evidenceKind':'APPLICATION_IMAGE_TELEMETRY','runtimeFacts':{'duration_seconds':47.3},'excerpt':'An image was generated.'}]
        self.assertTrue(history_provenance_failures('We implemented response caching [M21].',sources))
        self.assertFalse(history_provenance_failures('We updated the background [M21].',sources))
        self.assertTrue(history_provenance_failures('Added a faster parser. User-tested counts matched ten samples.\nAn image took 47.3 s [M21].',sources))
        sources.append({'id':'M22','evidenceKind':'USER_STATEMENT','excerpt':'I tested the updated parser and it returned the expected value.'})
        self.assertFalse(history_provenance_failures('The user tested the updated parser [M22].',sources))
        self.assertFalse(history_provenance_failures('Response caching was requested but is unverified [M21].',sources))
    def test_close_bridge_does_not_bypass_window_cleanup(self):
        from app import DesktopBridge
        fake=Mock();fake.window.close.return_value=False
        with patch('app.QApplication.instance') as instance:
            DesktopBridge.closeApp(fake)
        fake.window.close.assert_called_once();instance.assert_not_called()

if __name__=='__main__':unittest.main()
